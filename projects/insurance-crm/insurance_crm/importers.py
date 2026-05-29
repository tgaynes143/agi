"""Bulk import of companies + contacts from Excel/CSV.

The expected workbook format matches examples/sample_prospects.xlsx:
  Sheet "Companies":  name, domain, industry, revenue_band, employee_count,
                       hq_city, hq_state, lines_of_coverage, notes
  Sheet "Contacts":   company_domain, first_name, last_name, title, email,
                       email_confidence, linkedin_url, source, notes
"""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import load_workbook

from .db import DB, Company, Contact


def _row_to_dict(headers: list[str], row: tuple) -> dict:
    return {h: (v if v is not None else "") for h, v in zip(headers, row)}


def import_xlsx(db: DB, path: str | Path) -> dict:
    wb = load_workbook(Path(path), data_only=True)
    company_id_by_domain: dict[str, int] = {}
    companies_loaded = 0
    contacts_loaded = 0

    if "Companies" in wb.sheetnames:
        ws = wb["Companies"]
        rows = list(ws.iter_rows(values_only=True))
        if rows:
            headers = [str(h).strip() if h else "" for h in rows[0]]
            for row in rows[1:]:
                d = _row_to_dict(headers, row)
                if not d.get("name"):
                    continue
                domain = (d.get("domain") or "").strip().lower() or None
                cid = db.upsert_company(Company(
                    name=str(d["name"]).strip(),
                    domain=domain,
                    industry=d.get("industry") or None,
                    revenue_band=d.get("revenue_band") or None,
                    employee_count=int(d["employee_count"])
                        if d.get("employee_count") else None,
                    hq_city=d.get("hq_city") or None,
                    hq_state=d.get("hq_state") or None,
                    lines_of_coverage=d.get("lines_of_coverage") or None,
                    notes=d.get("notes") or None,
                ))
                companies_loaded += 1
                if domain:
                    company_id_by_domain[domain] = cid

    if "Contacts" in wb.sheetnames:
        ws = wb["Contacts"]
        rows = list(ws.iter_rows(values_only=True))
        if rows:
            headers = [str(h).strip() if h else "" for h in rows[0]]
            for row in rows[1:]:
                d = _row_to_dict(headers, row)
                if not (d.get("first_name") or d.get("email")):
                    continue
                domain = (d.get("company_domain") or "").strip().lower()
                company_id = company_id_by_domain.get(domain)
                if not company_id and domain:
                    # Company referenced but not in Companies sheet — create stub
                    company_id = db.upsert_company(Company(
                        name=domain, domain=domain,
                    ))
                    company_id_by_domain[domain] = company_id
                db.upsert_contact(Contact(
                    first_name=d.get("first_name") or None,
                    last_name=d.get("last_name") or None,
                    title=d.get("title") or None,
                    email=(d.get("email") or "").strip().lower() or None,
                    email_confidence=float(d["email_confidence"])
                        if d.get("email_confidence") not in (None, "") else 0.0,
                    linkedin_url=d.get("linkedin_url") or None,
                    source=d.get("source") or "import",
                    company_id=company_id,
                    notes=d.get("notes") or None,
                ))
                contacts_loaded += 1

    return {"companies": companies_loaded, "contacts": contacts_loaded}


def import_csv(db: DB, companies_csv: str | Path | None = None,
               contacts_csv: str | Path | None = None) -> dict:
    """Same idea, two separate CSVs."""
    companies_loaded = 0
    contacts_loaded = 0
    company_id_by_domain: dict[str, int] = {}

    if companies_csv:
        with open(companies_csv, newline="") as f:
            for d in csv.DictReader(f):
                if not d.get("name"):
                    continue
                domain = (d.get("domain") or "").strip().lower() or None
                cid = db.upsert_company(Company(
                    name=d["name"].strip(),
                    domain=domain,
                    industry=d.get("industry") or None,
                    revenue_band=d.get("revenue_band") or None,
                    employee_count=int(d["employee_count"])
                        if d.get("employee_count") else None,
                    hq_city=d.get("hq_city") or None,
                    hq_state=d.get("hq_state") or None,
                    lines_of_coverage=d.get("lines_of_coverage") or None,
                    notes=d.get("notes") or None,
                ))
                companies_loaded += 1
                if domain:
                    company_id_by_domain[domain] = cid

    if contacts_csv:
        with open(contacts_csv, newline="") as f:
            for d in csv.DictReader(f):
                if not (d.get("first_name") or d.get("email")):
                    continue
                domain = (d.get("company_domain") or "").strip().lower()
                company_id = company_id_by_domain.get(domain)
                if not company_id and domain:
                    company_id = db.upsert_company(Company(
                        name=domain, domain=domain,
                    ))
                    company_id_by_domain[domain] = company_id
                db.upsert_contact(Contact(
                    first_name=d.get("first_name") or None,
                    last_name=d.get("last_name") or None,
                    title=d.get("title") or None,
                    email=(d.get("email") or "").strip().lower() or None,
                    email_confidence=float(d["email_confidence"])
                        if d.get("email_confidence") else 0.0,
                    linkedin_url=d.get("linkedin_url") or None,
                    source=d.get("source") or "import",
                    company_id=company_id,
                    notes=d.get("notes") or None,
                ))
                contacts_loaded += 1

    return {"companies": companies_loaded, "contacts": contacts_loaded}
